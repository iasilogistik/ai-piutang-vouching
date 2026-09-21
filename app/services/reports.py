from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models import BillingReconciliation, ImportBatch, PhysicalBilling, SAPBilling, SPJ, VouchingResult
from app.services.control_evidence_dashboard import build_control_evidence_dashboard
from app.services.vouching import _net_document_amount

REPORT_ROOT = Path("storage/reports")


def _rows(db: Session, batch_id: int):
    batch = db.get(ImportBatch, batch_id)
    if not batch:
        raise ValueError("SAP import batch not found")
    stmt = (
        select(SAPBilling, BillingReconciliation, PhysicalBilling, VouchingResult)
        .join(BillingReconciliation, BillingReconciliation.sap_billing_id == SAPBilling.id, isouter=True)
        .join(PhysicalBilling, PhysicalBilling.id == BillingReconciliation.physical_billing_id, isouter=True)
        .join(VouchingResult, VouchingResult.billing_id == PhysicalBilling.id, isouter=True)
        .where(SAPBilling.import_batch_id == batch_id)
        .order_by(SAPBilling.id)
    )
    return batch, db.execute(stmt).all()


def _spj_partial(db: Session, physical: PhysicalBilling | None):
    if not physical or not physical.no_spj:
        return None
    matches = db.scalars(select(SPJ).where(SPJ.no_spj == physical.no_spj)).all()
    if len(matches) == 1:
        return matches[0].partial_payment
    return None


def _detail_values(db: Session, sap: SAPBilling, rec: BillingReconciliation | None, physical: PhysicalBilling | None, vouch: VouchingResult | None):
    billing_partial = physical.partial_payment if physical and physical.partial_payment is not None else 0
    spj_partial = _spj_partial(db, physical) or 0
    net_physical = _net_document_amount(physical.nominal if physical else None, billing_partial, spj_partial)
    return [
        sap.billing_document,
        sap.doc_date.isoformat(),
        float(sap.nominal),
        physical.billing_document if physical else "",
        physical.doc_date.isoformat() if physical and physical.doc_date else "",
        float(physical.nominal) if physical and physical.nominal is not None else "",
        float(billing_partial) if billing_partial else 0,
        float(spj_partial) if spj_partial else 0,
        float(net_physical) if net_physical is not None else "",
        float(rec.nominal_difference) if rec else float(sap.nominal),
        rec.status if rec else "NOT_FOUND",
        rec.exception_code if rec else "BILLING_DOCUMENT_NOT_FOUND",
        vouch.status if vouch else "NOT_FOUND",
        vouch.rule_code if vouch else "",
    ]


def _autosize(wb: Workbook) -> None:
    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 45)
            sheet.column_dimensions[column[0].column_letter].width = width


def build_report(db: Session, batch_id: int, fmt: str) -> Path:
    fmt = fmt.lower()
    if fmt not in {"xlsx", "pdf"}:
        raise ValueError("format must be xlsx or pdf")
    batch, rows = _rows(db, batch_id)
    REPORT_ROOT.mkdir(parents=True, exist_ok=True)
    stamp = datetime.utcnow().strftime("%Y%m%d%H%M%S")
    path = REPORT_ROOT / f"vouching_batch_{batch_id}_{stamp}.{fmt}"

    counts = {s: 0 for s in ("MATCH", "REVIEW", "EXCEPTION", "NOT_FOUND")}
    for _, rec, _, _ in rows:
        status = rec.status if rec else "NOT_FOUND"
        counts[status] = counts.get(status, 0) + 1

    headers = [
        "Billing Document SAP", "Doc Date SAP", "Nominal SAP", "Billing Document Fisik", "Doc Date Fisik",
        "Nominal Billing Fisik", "Partial Billing", "Partial SPJ", "Net Physical Nominal", "Difference",
        "Reconciliation", "Exception", "SPJ Result", "SPJ Rule",
    ]

    if fmt == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Summary"
        ws.append(["Batch ID", batch.id])
        ws.append(["File", batch.file_name])
        ws.append(["Period", batch.period.isoformat() if batch.period else ""])
        ws.append(["Total SAP records", len(rows)])
        for status in ("MATCH", "REVIEW", "EXCEPTION", "NOT_FOUND"):
            ws.append([status, counts[status]])
        detail = wb.create_sheet("Reconciliation")
        detail.append(headers)
        for sap, rec, physical, vouch in rows:
            detail.append(_detail_values(db, sap, rec, physical, vouch))
        _autosize(wb)
        wb.save(path)
    else:
        styles = getSampleStyleSheet()
        doc = SimpleDocTemplate(str(path), pagesize=landscape(A4), rightMargin=24, leftMargin=24, topMargin=24, bottomMargin=24)
        story = [Paragraph(f"AI Piutang Vouching — Batch {batch_id}", styles["Title"]), Spacer(1, 8),
                 Paragraph(f"File: {batch.file_name} | Total SAP records: {len(rows)}", styles["Normal"]), Spacer(1, 10)]
        summary = [["MATCH", counts["MATCH"]], ["REVIEW", counts["REVIEW"]], ["EXCEPTION", counts["EXCEPTION"]], ["NOT_FOUND", counts["NOT_FOUND"]]]
        table = Table([["Status", "Count"]] + summary, colWidths=[120, 80])
        table.setStyle(TableStyle([("GRID", (0,0), (-1,-1), 0.5, colors.black), ("BACKGROUND", (0,0), (-1,0), colors.lightgrey)]))
        story += [table, Spacer(1, 14)]
        data = [["Billing SAP", "Nominal SAP", "Nominal Billing", "Partial Billing", "Partial SPJ", "Net Physical", "Selisih", "Reconcile", "Exception", "SPJ"]]
        for sap, rec, physical, vouch in rows:
            values = _detail_values(db, sap, rec, physical, vouch)
            data.append([values[0], values[2], values[5], values[6], values[7], values[8], values[9], values[10], values[11], values[12]])
        detail = Table(data, repeatRows=1)
        detail.setStyle(TableStyle([("GRID", (0,0), (-1,-1), 0.35, colors.black), ("BACKGROUND", (0,0), (-1,0), colors.lightgrey), ("FONTSIZE", (0,0), (-1,-1), 7)]))
        story.append(detail)
        doc.build(story)
    return path


def build_control_evidence_report(db: Session, *, review_only: bool = False, limit: int = 500) -> Path:
    REPORT_ROOT.mkdir(parents=True, exist_ok=True)
    stamp = datetime.utcnow().strftime("%Y%m%d%H%M%S")
    path = REPORT_ROOT / f"spj_control_evidence_{stamp}.xlsx"
    dashboard = build_control_evidence_dashboard(db, review_only=review_only, limit=limit)

    wb = Workbook()
    summary_sheet = wb.active
    summary_sheet.title = "Summary"
    summary_sheet.append(["Metric", "Value"])
    summary_sheet.append(["Total SPJ Evidence", dashboard["summary"]["total_documents"]])
    summary_sheet.append(["PASS", dashboard["summary"].get("overall_control_status", {}).get("PASS", 0)])
    summary_sheet.append(["REVIEW", dashboard["summary"].get("overall_control_status", {}).get("REVIEW", 0)])
    summary_sheet.append(["EXCEPTION", dashboard["summary"].get("overall_control_status", {}).get("EXCEPTION", 0)])
    summary_sheet.append(["Review Required", dashboard["summary"]["review_required_documents"]])

    rows_sheet = wb.create_sheet("Control Evidence")
    headers = [
        "Control Evidence ID", "Document ID", "File Name", "No SPJ", "Billing ID", "SPJ Vouching Status",
        "Overall Control Status", "Review Required", "Review Status", "Reviewer", "Reviewer Remarks",
        "TTD Penerima", "TTD Driver", "TTD Satpam", "TTD BM", "TTD Checker", "Stempel", "Nama Stempel",
        "Stempel vs Customer", "Alasan Review", "Document URL",
    ]
    rows_sheet.append(headers)
    for row in dashboard["rows"]:
        rows_sheet.append([
            row.get("control_evidence_id"),
            row.get("document_id"),
            row.get("file_name"),
            row.get("no_spj"),
            row.get("billing_id"),
            row.get("spj_vouching_status"),
            row.get("overall_control_status"),
            row.get("review_required"),
            row.get("review_status"),
            row.get("reviewer_id"),
            row.get("reviewer_remarks"),
            row.get("receiver_signature_status"),
            row.get("driver_signature_status"),
            row.get("security_signature_status"),
            row.get("bm_signature_status"),
            row.get("checker_signature_status"),
            row.get("receiver_stamp_status"),
            row.get("stamp_text_raw"),
            row.get("stamp_customer_match_status"),
            "; ".join(row.get("review_reasons") or []),
            row.get("document_url"),
        ])

    queue_sheet = wb.create_sheet("Manual Review Queue")
    queue_sheet.append(headers)
    for row in dashboard["manual_review_queue"]:
        queue_sheet.append([
            row.get("control_evidence_id"), row.get("document_id"), row.get("file_name"), row.get("no_spj"),
            row.get("billing_id"), row.get("spj_vouching_status"), row.get("overall_control_status"),
            row.get("review_required"), row.get("review_status"), row.get("reviewer_id"), row.get("reviewer_remarks"),
            row.get("receiver_signature_status"), row.get("driver_signature_status"), row.get("security_signature_status"),
            row.get("bm_signature_status"), row.get("checker_signature_status"), row.get("receiver_stamp_status"),
            row.get("stamp_text_raw"), row.get("stamp_customer_match_status"), "; ".join(row.get("review_reasons") or []),
            row.get("document_url"),
        ])

    _autosize(wb)
    wb.save(path)
    return path

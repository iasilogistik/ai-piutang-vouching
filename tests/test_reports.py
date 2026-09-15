from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.models import BillingReconciliation, ImportBatch, SAPBilling
from app.services.reports import build_report


def test_build_xlsx_report(tmp_path, monkeypatch):
    engine = create_engine("sqlite:///:memory:")
    from app.database import Base
    Base.metadata.create_all(engine)
    import app.services.reports as reports
    monkeypatch.setattr(reports, "REPORT_ROOT", tmp_path)
    with Session(engine) as db:
        batch = ImportBatch(file_name="sap.xlsx", period=date(2026, 9, 1), total_records=1, status="IMPORTED")
        db.add(batch); db.flush()
        sap = SAPBilling(import_batch_id=batch.id, billing_document="B-001", doc_date=date(2026, 9, 1), nominal=Decimal("100.00"))
        db.add(sap); db.flush()
        db.add(BillingReconciliation(sap_billing_id=sap.id, billing_match=False, date_match=False, nominal_match=False,
                                     nominal_difference=Decimal("100.00"), status="EXCEPTION", exception_code="BILLING_DOCUMENT_NOT_FOUND"))
        db.commit()
        path = build_report(db, batch.id, "xlsx")
        assert path.exists()
        wb = load_workbook(path)
        assert "Summary" in wb.sheetnames
        assert "Reconciliation" in wb.sheetnames
        assert wb["Summary"]["A4"].value == "Total SAP records"


def test_build_pdf_report(tmp_path, monkeypatch):
    engine = create_engine("sqlite:///:memory:")
    from app.database import Base
    Base.metadata.create_all(engine)
    import app.services.reports as reports
    monkeypatch.setattr(reports, "REPORT_ROOT", tmp_path)
    with Session(engine) as db:
        batch = ImportBatch(file_name="sap.xlsx", total_records=0, status="IMPORTED")
        db.add(batch); db.commit()
        path = build_report(db, batch.id, "pdf")
        assert path.exists()
        assert path.suffix == ".pdf"
        assert path.stat().st_size > 0

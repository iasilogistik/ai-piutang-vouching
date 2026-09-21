from __future__ import annotations

from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib import colors

from app.models import AuditReport

REPORT_ROOT = Path("storage/reports")


def _safe(value: str) -> str:
    return "".join(ch if ch.isalnum() or ch in {"-", "_"} else "_" for ch in value).strip("_") or "branch"


def _summary_rows(report: AuditReport) -> list[list[object]]:
    return [
        ["Report ID", report.id],
        ["Branch", report.branch],
        ["Period Start", report.period_start.isoformat()],
        ["Period End", report.period_end.isoformat()],
        ["Status", report.status],
        ["Population", report.population_count],
        ["Sampled", report.sampled_count],
        ["Matched", report.matched_count],
        ["Exceptions", report.exception_count],
        ["Unresolved Exceptions", report.unresolved_exception_count],
        ["Resolved Exceptions", report.resolved_exception_count],
        ["Created By", report.created_by or ""],
        ["Approved By", report.approved_by or ""],
        ["Approved At", report.approved_at.isoformat() if report.approved_at else ""],
    ]


def build_audit_report_export(report: AuditReport, fmt: str) -> Path:
    fmt = fmt.strip().lower()
    if fmt not in {"xlsx", "pdf"}:
        raise ValueError("format must be xlsx or pdf")

    REPORT_ROOT.mkdir(parents=True, exist_ok=True)
    stamp = datetime.utcnow().strftime("%Y%m%d%H%M%S")
    path = REPORT_ROOT / f"audit_report_{report.id}_{_safe(report.branch)}_{stamp}.{fmt}"

    if fmt == "xlsx":
        wb = Workbook()
        ws = wb.active
        ws.title = "Audit Report"
        ws.append(["AI Piutang Vouching - Audit Report"])
        ws.append([])
        for row in _summary_rows(report):
            ws.append(row)
        ws.append([])
        ws.append(["Finding Summary"])
        ws.append([report.finding_summary or ""])
        ws.append([])
        ws.append(["Conclusion"])
        ws.append([report.conclusion or ""])

        ws.freeze_panes = "A3"
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 48
        ws["A1"].font = ws["A1"].font.copy(bold=True, size=14)
        for cell in ws["A"]:
            cell.alignment = cell.alignment.copy(vertical="top", wrap_text=True)
        for cell in ws["B"]:
            cell.alignment = cell.alignment.copy(vertical="top", wrap_text=True)
        wb.save(path)
        return path

    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )
    story = [
        Paragraph("AI Piutang Vouching - Audit Report", styles["Title"]),
        Spacer(1, 12),
    ]
    summary = Table(_summary_rows(report), colWidths=[150, 330])
    summary.setStyle(
        TableStyle(
            [
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("BACKGROUND", (0, 0), (0, -1), colors.whitesmoke),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
            ]
        )
    )
    story.extend(
        [
            summary,
            Spacer(1, 16),
            Paragraph("Finding Summary", styles["Heading2"]),
            Paragraph(escape(report.finding_summary or "-"), styles["BodyText"]),
            Spacer(1, 12),
            Paragraph("Conclusion", styles["Heading2"]),
            Paragraph(escape(report.conclusion or "-"), styles["BodyText"]),
        ]
    )
    doc.build(story)
    return path

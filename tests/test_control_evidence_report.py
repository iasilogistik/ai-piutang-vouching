from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.database import Base
from app.models import Document, DocumentControlEvidence, SPJ
from app.services.reports import build_control_evidence_report


def _db():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    return engine


def test_build_control_evidence_report_creates_workbook():
    with Session(_db()) as db:
        doc = Document(
            file_name="spj-report.pdf",
            file_type="PDF",
            document_type="SPJ",
            file_hash="evidence-report",
            storage_path="storage/spj-report.pdf",
        )
        db.add(doc)
        db.flush()
        db.add(SPJ(document_id=doc.id, no_spj_raw="SPJ-100", no_spj="SPJ100"))
        db.add(
            DocumentControlEvidence(
                document_id=doc.id,
                receiver_signature_status="PRESENT",
                driver_signature_status="PRESENT",
                security_signature_status="UNKNOWN",
                bm_signature_status="PRESENT",
                checker_signature_status="PRESENT",
                receiver_stamp_status="PRESENT",
                stamp_text_raw="TOKO SANTOSO",
                stamp_customer_match_status="MATCH",
                review_required=True,
                review_reasons="Satpam perlu dicek manual",
            )
        )
        db.commit()

        path = build_control_evidence_report(db, review_only=True)

        assert path.exists()
        workbook = load_workbook(path)
        assert "Summary" in workbook.sheetnames
        assert "Control Evidence" in workbook.sheetnames
        assert "Manual Review Queue" in workbook.sheetnames
        rows = list(workbook["Control Evidence"].iter_rows(values_only=True))
        assert rows[0][0] == "Control Evidence ID"
        assert rows[1][3] == "SPJ100"
        assert rows[1][18] == "MATCH"
